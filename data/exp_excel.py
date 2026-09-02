import os
import openpyxl
from openpyxl.styles import Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from datetime import date
from tkinter import messagebox as mb
from tkinter import filedialog

def exportar_excel(arc, lista, datos):

    pathinit = 'exportar/reportes/'
    file_path = filedialog.askdirectory(initialdir=pathinit)
    if file_path:
        today = date.today()
        pdf_name = arc+"_"+today.strftime("%d_%m_%Y")+".xlsx"
        xlsSave = os.path.join(file_path, pdf_name)   

        excelWorkbook = openpyxl.Workbook()
        excelWorksheet = excelWorkbook.active
        border_style = Border(left=Side(border_style='thin'), 
                            right=Side(border_style='thin'), 
                            top=Side(border_style='thin'), 
                            bottom=Side(border_style='thin'))

        for i, val in enumerate(lista, start=1):
            excelWorksheet.cell(row=1, column=i).value = val
            cell = excelWorksheet.cell(row=1, column=i)
            cell.value = val
            cell.border = border_style
        
        fila = 2
        try:
            for j in datos:
                j
                for i in range(len(lista)):
                    # excelWorksheet.cell(row=fila, column=1+i).value = values[i]
                    cell = excelWorksheet.cell(row=fila, column=1+i)
                    cell.value = j[i]
                fila += 1
            table = Table(displayName="Table1", ref=excelWorksheet.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                                                showLastColumn=False, showRowStripes=False, showColumnStripes=False)
            excelWorksheet.add_table(table)
            
            # Ajustar el ancho de las columnas automáticamente
            for column_cells in excelWorksheet.columns:
                length = max(len(str(cell.value)) for cell in column_cells)
                excelWorksheet.column_dimensions[column_cells[0].column_letter].width = length + 2
            excelWorkbook.save(xlsSave)

            mb.showinfo("Exportar Datos", "Los datos se han exportado correctamente.")
        except Exception as e:
            mb.showerror(message=f"Error al exportar, Error: {str(e)}", title="¡Atención!")
            print(e)